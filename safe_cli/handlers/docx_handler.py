from pathlib import Path
from typing import Dict, Optional

from docx import Document

from ..column_rules import placeholder_for_header
from ..engine import AnonymizerCore


def _anonymize_paragraph(
    para,
    engine: AnonymizerCore,
    stats: Dict[str, int],
    forced: Optional[str] = None,
) -> None:
    full_text = para.text
    if not full_text.strip():
        return

    if forced:
        anonymized = forced
        stats[forced] = stats.get(forced, 0) + 1
    else:
        anonymized, para_stats = engine.anonymize_text(full_text)
        for k, v in para_stats.items():
            stats[k] = stats.get(k, 0) + v

    if anonymized == full_text:
        return

    runs = para.runs
    if runs:
        for run in runs[1:]:
            run._r.getparent().remove(run._r)
        runs[0].text = anonymized
    else:
        para.add_run(anonymized)


def _process_table(table, engine: AnonymizerCore, stats: Dict[str, int]) -> None:
    """Process a table, using the first row as column headers for forced anonymization."""
    rows = table.rows
    if not rows:
        return

    # Build forced-placeholder map from first row (header row)
    col_forced: Dict[int, Optional[str]] = {}
    for col_idx, cell in enumerate(rows[0].cells):
        header = cell.text.strip()
        col_forced[col_idx] = placeholder_for_header(header)

    for row_idx, row in enumerate(rows):
        for col_idx, cell in enumerate(row.cells):
            forced = col_forced.get(col_idx) if row_idx > 0 else None
            for para in cell.paragraphs:
                _anonymize_paragraph(para, engine, stats, forced=forced)


def process(input_path: Path, output_path: Path, engine: AnonymizerCore) -> Dict[str, int]:
    stats: Dict[str, int] = {}
    doc = Document(str(input_path))

    for para in doc.paragraphs:
        _anonymize_paragraph(para, engine, stats)

    for table in doc.tables:
        _process_table(table, engine, stats)

    for section in doc.sections:
        for hf in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            if hf is not None:
                for para in hf.paragraphs:
                    _anonymize_paragraph(para, engine, stats)

    doc.save(str(output_path))
    return stats
