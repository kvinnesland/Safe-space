from pathlib import Path
from typing import Dict, Optional

import openpyxl

from ..column_rules import placeholder_for_header
from ..engine import AnonymizerCore
from ..media import PLACEHOLDER_IMAGE


def _remove_images(sheet, stats: Dict[str, int]) -> None:
    """Remove all embedded images from a worksheet."""
    if hasattr(sheet, "_images") and sheet._images:
        count = len(sheet._images)
        sheet._images.clear()
        stats[PLACEHOLDER_IMAGE] = stats.get(PLACEHOLDER_IMAGE, 0) + count


def process(input_path: Path, output_path: Path, engine: AnonymizerCore) -> Dict[str, int]:
    stats: Dict[str, int] = {}
    wb = openpyxl.load_workbook(input_path)

    for sheet in wb.worksheets:
        _remove_images(sheet, stats)

        col_forced: Dict[int, str] = {}
        col_headers: Dict[int, str] = {}

        for cell in next(sheet.iter_rows(min_row=1, max_row=1), []):
            if not isinstance(cell.value, str) or not cell.value.strip():
                continue
            header = cell.value.strip()
            forced = placeholder_for_header(header)
            if forced:
                col_forced[cell.column] = forced
            col_headers[cell.column] = header.lower()

        for row_idx, row in enumerate(sheet.iter_rows(), start=1):
            for cell in row:
                if not isinstance(cell.value, str) or not cell.value.strip():
                    continue

                if row_idx == 1:
                    continue  # never anonymize the header row itself

                value = cell.value

                forced = col_forced.get(cell.column)
                if forced:
                    if value != forced:
                        stats[forced] = stats.get(forced, 0) + 1
                        cell.value = forced
                    continue

                header_text = col_headers.get(cell.column, "")
                if header_text:
                    text = f"{header_text}: {value}"
                    anon_text, cell_stats = engine.anonymize_text(text)
                    prefix = f"{header_text}: "
                    if anon_text.startswith(prefix):
                        anonymized = anon_text[len(prefix):]
                    else:
                        anonymized, cell_stats = engine.anonymize_text(value)
                else:
                    anonymized, cell_stats = engine.anonymize_text(value)

                for k, v in cell_stats.items():
                    stats[k] = stats.get(k, 0) + v
                if anonymized != value:
                    cell.value = anonymized

    wb.save(output_path)
    return stats
